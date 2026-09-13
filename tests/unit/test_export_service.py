import pytest
import json
import csv
from pathlib import Path
from src.core.dtos.export_dtos import ExportConfigDTO, ExportFormat
from src.services.export_service import ExportService, OPENPYXL_AVAILABLE

if OPENPYXL_AVAILABLE:
    import openpyxl

class MockCommentRepo:
    def get_comments_for_drawing(self, drawing_id):
        return [
            {
                'id': '1',
                'drawing_id': drawing_id,
                'status': 'Pending',
                'raw_text': 'Pipe clearance less than 50mm from structural beam',
                'category_name': 'Coordination/Interference',
                'confidence': 0.92,
                'page_number': 1,
            },
            {
                'id': '2',
                'drawing_id': drawing_id,
                'status': 'Approved',
                'raw_text': 'Dimension missing on flange weld neck',
                'category_name': 'Dimensional/Tolerancing',
                'confidence': 0.97,
                'page_number': 2,
            },
        ]

class MockProjectRepo:
    pass

class MockDrawingRepo:
    def get_drawing_by_id(self, drawing_id):
        return {
            "id": drawing_id,
            "file_name": "M-70086-01-013_B_.pdf",
            "title": "Primary Crusher Piping Isometric",
            "author": "J. Doe",
        }

def test_export_to_json_creates_file(tmp_path):
    repo = MockCommentRepo()
    service = ExportService(repo, MockProjectRepo(), MockDrawingRepo())
    out_file = tmp_path / "out.json"
    
    config = ExportConfigDTO(
        output_path=out_file,
        format=ExportFormat.JSON,
        drawing_id="draw1",
        contract_no="CTR-2026-01",
        plant_name="Austin Substation",
        epod_wo_no="WO-440192",
        designer_name="Lead Reviewer"
    )
    res = service.export_drawing_comments(config)
    
    assert res.success is True
    assert res.total_rows == 2
    assert out_file.exists()
    
    with open(out_file, encoding='utf-8') as f:
        data = json.load(f)
        assert data['metadata']['total_records'] == 2
        assert data['metadata']['contract_no'] == "CTR-2026-01"
        assert len(data['comments']) == 2

def test_export_to_csv_creates_file(tmp_path):
    repo = MockCommentRepo()
    service = ExportService(repo, MockProjectRepo(), MockDrawingRepo())
    out_file = tmp_path / "out.csv"
    
    config = ExportConfigDTO(
        output_path=out_file,
        format=ExportFormat.CSV,
        drawing_id="draw1",
        contract_no="CTR-2026-01",
        plant_name="Austin Substation",
        epod_wo_no="WO-440192",
        designer_name="Lead Reviewer",
        department_name="Piping Engineering",
    )
    res = service.export_drawing_comments(config)
    
    assert res.success is True
    assert res.total_rows == 2
    assert out_file.exists()
    
    with open(out_file, newline='', encoding='utf-8') as f:
        reader = csv.reader(f)
        rows = list(reader)
        # Header + 2 data rows
        assert len(rows) == 3
        assert len(rows[0]) == 10
        assert rows[0][0] == "Date"
        assert rows[0][1] == "Contract #"
        assert rows[0][7] == "Errors Description"
        assert rows[0][8] == "Category of Error"
        assert rows[0][9] == "Engineering Department"
        # Check data row 1
        assert rows[1][1] == "CTR-2026-01"
        assert rows[1][2] == "Austin Substation"
        assert rows[1][7] == "Pipe clearance less than 50mm from structural beam"
        assert rows[1][8] == "Coordination/Interference"
        assert rows[1][9] == "Piping Engineering"

def test_export_result_dto_fields(tmp_path):
    repo = MockCommentRepo()
    service = ExportService(repo, MockProjectRepo(), MockDrawingRepo())
    out_file = tmp_path / "out.json"
    
    config = ExportConfigDTO(output_path=out_file, format=ExportFormat.JSON, drawing_id="draw1")
    res = service.export_drawing_comments(config)
    
    assert hasattr(res, 'output_path')
    assert hasattr(res, 'format')
    assert hasattr(res, 'total_rows')
    assert hasattr(res, 'total_sheets')
    assert hasattr(res, 'file_size_bytes')
    assert hasattr(res, 'success')
    assert hasattr(res, 'error_message')

def test_export_config_defaults():
    config = ExportConfigDTO(output_path=Path("test.xlsx"))
    assert config.format == 'xlsx'
    assert config.include_summary_sheet is True
    assert config.include_confidence_scores is True
    assert config.filter_status is None
    assert config.drawing_id is None
    assert config.department_name is None

@pytest.mark.skipif(not OPENPYXL_AVAILABLE, reason="openpyxl not installed")
def test_error_tracker_excel_structure_and_styling(tmp_path):
    repo = MockCommentRepo()
    dwg_repo = MockDrawingRepo()
    service = ExportService(repo, MockProjectRepo(), dwg_repo)
    out_file = tmp_path / "error_tracker_test.xlsx"
    
    config = ExportConfigDTO(
        output_path=out_file,
        format=ExportFormat.EXCEL,
        drawing_id="draw1",
        date_str="2026-09-12",
        contract_no="CTR-2026-881",
        plant_name="Austin Substation",
        epod_wo_no="WO-99014",
        designer_name="Soham Lead",
        drawing_no="M-70086-01-013_B_",
        drawing_title="Primary Crusher Piping Isometric",
        department_name="Piping Engineering",
    )
    res = service.export_drawing_comments(config)
    
    assert res.success is True
    assert out_file.exists()
    assert res.file_size_bytes > 0
    assert res.total_rows == 2
    # Executive Summary + 7 standard UCC departments
    assert res.total_sheets == 8

    # Verify Excel internal cell structure and headers
    wb = openpyxl.load_workbook(out_file)

    # 1. Verify sheet names
    expected_sheets = [
        "Executive Summary",
        "Electrical Engineering",
        "GPD",
        "Pipe Support Engineering",
        "Piping Engineering",
        "Plakon",
        "Structural & Physical Design",
        "System Engineering",
    ]
    assert wb.sheetnames == expected_sheets

    # 2. Check "Piping Engineering" worksheet (where our mock comments belong)
    ws = wb["Piping Engineering"]

    # Check Row 1 (Group Banner tier - 9 columns)
    assert ws["A1"].value == "Standard Input Field"
    assert ws["C1"].value == "Auto Read by Program"
    assert ws["D1"].value == "Standard Input Field"
    assert ws["F1"].value == "Auto Read by Program"

    # Check Fill Colors (Orange = FFC000, Green = 92D050)
    assert ws["A1"].fill.start_color.rgb in ("00FFC000", "FFC000")
    assert ws["C1"].fill.start_color.rgb in ("0092D050", "92D050")

    # Check Row 2 (Column Number tier: 1, 2, 3, 4, 6, 7, 8, 9, 10)
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == 2
    assert ws.cell(row=2, column=3).value == 3
    assert ws.cell(row=2, column=4).value == 4
    assert ws.cell(row=2, column=5).value == 6  # Column 6 matching template
    assert ws.cell(row=2, column=6).value == 7
    assert ws.cell(row=2, column=7).value == 8
    assert ws.cell(row=2, column=8).value == 9
    assert ws.cell(row=2, column=9).value == 10

    # Check Row 3 (Source tier - 9 columns)
    assert ws.cell(row=3, column=1).value == "User Input"
    assert "Drawing #" in ws.cell(row=3, column=6).value
    assert "Commentary" in ws.cell(row=3, column=8).value
    assert "Classify Error" in ws.cell(row=3, column=9).value

    # Check Row 4 (Primary Column Headers - 9 columns)
    expected_headers = [
        "Date", "Contract #", "Plant Name", "E-Pod WO #", "UCC-I Designer",
        "Drawing #", "Drawing Title", "Errors Description", "Category of Error"
    ]
    for idx, expected in enumerate(expected_headers, 1):
        assert ws.cell(row=4, column=idx).value == expected

    # Check Row 5 (Data Row 1)
    assert ws.cell(row=5, column=1).value == "2026-09-12"
    assert ws.cell(row=5, column=2).value == "CTR-2026-881"
    assert ws.cell(row=5, column=3).value == "Austin Substation"
    assert ws.cell(row=5, column=4).value == "WO-99014"
    assert ws.cell(row=5, column=5).value == "Soham Lead"
    assert ws.cell(row=5, column=6).value == "M-70086-01-013_B_"
    assert ws.cell(row=5, column=7).value == "Primary Crusher Piping Isometric"
    assert ws.cell(row=5, column=8).value == "Pipe clearance less than 50mm from structural beam"
    assert ws.cell(row=5, column=9).value == "Coordination/Interference"

    # Check Row 6 (Data Row 2)
    assert ws.cell(row=6, column=8).value == "Dimension missing on flange weld neck"
    assert ws.cell(row=6, column=9).value == "Dimensional/Tolerancing"

    # 3. Check an empty department sheet (e.g. "Electrical Engineering") has 5 empty placeholder rows
    ws_elec = wb["Electrical Engineering"]
    assert ws_elec.cell(row=4, column=1).value == "Date"
    # Row 5 to Row 9 should be placeholder rows with empty string values
    for r in range(5, 10):
        for col_idx in range(1, 10):
            assert ws_elec.cell(row=r, column=col_idx).value in ("", None)



@pytest.mark.skipif(not OPENPYXL_AVAILABLE, reason="openpyxl not installed")
def test_unassigned_department_tab_creation(tmp_path):
    class MockUnassignedRepo:
        def get_comments_for_drawing(self, drawing_id):
            return [
                {
                    'id': '1',
                    'drawing_id': drawing_id,
                    'status': 'Pending',
                    'raw_text': 'Cable tray clearance issue',
                    'category_name': 'Electrical',
                    'department_name': 'Electrical Engineering',
                },
                {
                    'id': '2',
                    'drawing_id': drawing_id,
                    'status': 'Pending',
                    'raw_text': 'Unknown issue with no department assigned',
                    'category_name': 'General',
                    'department_name': 'Unassigned',
                },
            ]

    service = ExportService(MockUnassignedRepo(), MockProjectRepo(), MockDrawingRepo())
    out_file = tmp_path / "unassigned_test.xlsx"

    config = ExportConfigDTO(
        output_path=out_file,
        format=ExportFormat.EXCEL,
        drawing_id="draw1",
        include_summary_sheet=False,
    )
    res = service.export_drawing_comments(config)

    assert res.success is True
    wb = openpyxl.load_workbook(out_file)
    assert "Unassigned" in wb.sheetnames

    ws_unassigned = wb["Unassigned"]
    assert ws_unassigned.cell(row=5, column=8).value == "Unknown issue with no department assigned"

