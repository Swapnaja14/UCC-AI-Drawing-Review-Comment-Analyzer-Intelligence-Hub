"""
tests/unit/test_department_upload_workflow.py
Unit tests verifying Department-Wise Drawing Upload workflow, database persistence on DrawingModel & CommentModel, multi-sheet Excel export segregation per department, and department-filtered history & analytics.
"""

import pytest
from pathlib import Path
import fitz  # PyMuPDF
from datetime import datetime, timezone

from src.services.file_service import FileService
from src.services.pdf_service import PDFService
from src.services.workflow_engine import ProcessingWorkflowEngine
from src.services.export_service import ExportService
from src.services.analytics_service import AnalyticsService
from src.infrastructure.pdf.pymupdf_adapter import PyMuPDFAdapter
from src.infrastructure.storage.repository import (
    DatabaseEngine,
    DrawingRepository,
    CommentRepository,
    EngineeringDepartmentRepository,
    ProjectRepository,
)
from src.infrastructure.storage.models import DrawingModel, CommentModel, EngineeringDepartmentModel
from src.core.dtos.workflow_dtos import WorkflowState
from src.core.dtos.export_dtos import ExportConfigDTO, ExportFormat


@pytest.fixture
def test_db_engine(tmp_path: Path) -> DatabaseEngine:
    db_file = tmp_path / "test_dept_workflow.db"
    return DatabaseEngine(db_path=db_file)


@pytest.fixture
def sample_pdf_drawing(tmp_path: Path) -> Path:
    pdf_file = tmp_path / "piping_test_drawing.pdf"
    doc = fitz.open()
    page = doc.new_page(width=841.89, height=595.28)
    page.insert_text((50, 50), "DRAWING NO: UCC-DWG-PIPING-101", fontsize=14)
    page.insert_text((50, 80), "PIPING SYSTEM DIAGRAM", fontsize=12)
    doc.save(pdf_file)
    doc.close()
    return pdf_file


def test_department_repository_seed_and_all_departments(test_db_engine: DatabaseEngine):
    dept_repo = EngineeringDepartmentRepository(test_db_engine)
    depts = dept_repo.get_all_departments()

    assert len(depts) >= 7
    dept_names = [d["name"] for d in depts]
    assert "Electrical Engineering" in dept_names
    assert "Piping Engineering" in dept_names
    assert "GPD" in dept_names
    assert "Structural & Physical Design" in dept_names


def test_drawing_repository_save_with_department(test_db_engine: DatabaseEngine, sample_pdf_drawing: Path):
    pdf_service = PDFService(pdf_loader=PyMuPDFAdapter())
    doc_dto = pdf_service.process_pdf_document(sample_pdf_drawing)
    drawing_repo = DrawingRepository(test_db_engine)
    dept_repo = EngineeringDepartmentRepository(test_db_engine)

    piping_dept = dept_repo.get_or_create_department("Piping Engineering")
    dept_id = piping_dept["id"]

    record = drawing_repo.save_drawing_from_dto(doc_dto, department_id=dept_id)

    assert record["id"].startswith("DWG-")
    assert record["department_id"] == dept_id

    # Verify query directly from session
    with test_db_engine.get_session() as session:
        dwg_model = session.get(DrawingModel, record["id"])
        assert dwg_model is not None
        assert dwg_model.department_id == dept_id
        assert dwg_model.department_rel.name == "Piping Engineering"


def test_workflow_engine_execution_with_department(test_db_engine: DatabaseEngine, sample_pdf_drawing: Path):
    drawing_repo = DrawingRepository(test_db_engine)
    comment_repo = CommentRepository(test_db_engine)
    dept_repo = EngineeringDepartmentRepository(test_db_engine)

    electrical_dept = dept_repo.get_or_create_department("Electrical Engineering")
    dept_id = electrical_dept["id"]

    file_service = FileService()
    pdf_service = PDFService(pdf_loader=PyMuPDFAdapter())

    engine = ProcessingWorkflowEngine(
        file_service=file_service,
        pdf_service=pdf_service,
        drawing_repo=drawing_repo,
        comment_repo=comment_repo,
    )

    steps_recorded = []
    def on_step(snapshot):
        steps_recorded.append(snapshot)

    result = engine.execute_workflow(
        sample_pdf_drawing,
        progress_callback=on_step,
        department_id=dept_id,
    )

    assert result.status == "Completed"
    assert result.file_name == "piping_test_drawing.pdf"
    assert engine.current_state == WorkflowState.COMPLETED

    # Verify drawing department_id in DB
    dwg = drawing_repo.get_drawing_by_id(result.drawing_id)
    assert dwg is not None
    assert dwg["department_id"] == dept_id


def test_comment_inherits_drawing_department(test_db_engine: DatabaseEngine, sample_pdf_drawing: Path):
    pdf_service = PDFService(pdf_loader=PyMuPDFAdapter())
    doc_dto = pdf_service.process_pdf_document(sample_pdf_drawing)
    drawing_repo = DrawingRepository(test_db_engine)
    comment_repo = CommentRepository(test_db_engine)
    dept_repo = EngineeringDepartmentRepository(test_db_engine)

    gpd_dept = dept_repo.get_or_create_department("GPD")
    dept_id = gpd_dept["id"]

    dwg_rec = drawing_repo.save_drawing_from_dto(doc_dto, department_id=dept_id)
    drawing_id = dwg_rec["id"]

    # Create dummy comment
    cmt_id = comment_repo.save_comment(
        drawing_id=drawing_id,
        page_number=1,
        raw_text="Check pipe support clearance",
        cleaned_text="Check pipe support clearance",
        bbox=(10.0, 10.0, 50.0, 50.0),
        confidence=0.92,
        category_name="Dimensional",
        department_id=dept_id,
    )

    assert cmt_id["id"].startswith("CMT-")

    comments = comment_repo.get_comments_for_drawing(drawing_id)
    assert len(comments) == 1
    assert comments[0]["department_id"] == dept_id
    assert comments[0]["department_name"] == "GPD"


def test_excel_export_multi_sheet_department_segregation(test_db_engine: DatabaseEngine, tmp_path: Path, sample_pdf_drawing: Path):
    pdf_service = PDFService(pdf_loader=PyMuPDFAdapter())
    doc_dto = pdf_service.process_pdf_document(sample_pdf_drawing)
    comment_repo = CommentRepository(test_db_engine)
    dept_repo = EngineeringDepartmentRepository(test_db_engine)
    proj_repo = ProjectRepository(test_db_engine)
    dwg_repo = DrawingRepository(test_db_engine)

    piping_dept = dept_repo.get_or_create_department("Piping Engineering")
    electrical_dept = dept_repo.get_or_create_department("Electrical Engineering")

    dwg_piping = dwg_repo.save_drawing_from_dto(doc_dto, department_id=piping_dept["id"])
    
    import dataclasses
    doc_dto_elec = dataclasses.replace(pdf_service.process_pdf_document(sample_pdf_drawing), file_hash_sha256="elechash999")
    dwg_elec = dwg_repo.save_drawing_from_dto(doc_dto_elec, department_id=electrical_dept["id"])

    # Add comments for Piping
    comment_repo.save_comment(
        drawing_id=dwg_piping["id"],
        page_number=1,
        raw_text="Piping clash with tray",
        bbox=(0, 0, 10, 10),
        category_name="Technical",
        department_id=piping_dept["id"],
    )

    # Add comments for Electrical
    comment_repo.save_comment(
        drawing_id=dwg_elec["id"],
        page_number=1,
        raw_text="Electrical conduit ungrounded",
        bbox=(0, 0, 10, 10),
        category_name="Standards",
        department_id=electrical_dept["id"],
    )

    export_service = ExportService(
        comment_repo=comment_repo,
        project_repo=proj_repo,
        drawing_repo=dwg_repo,
        department_repo=dept_repo,
    )

    out_xlsx = tmp_path / "UCC_Multi_Dept_Error_Tracker.xlsx"
    config = ExportConfigDTO(
        output_path=str(out_xlsx),
        format=ExportFormat.EXCEL,
        drawing_id=dwg_piping["id"],
    )

    # Re-query all comments
    comments = comment_repo.get_comments_for_drawing(dwg_piping["id"]) + comment_repo.get_comments_for_drawing(dwg_elec["id"])
    result = export_service._export_to_excel(comments, config)

    assert result.success is True
    assert result.total_sheets >= 7  # All active UCC departments generated

    # Read back workbook using openpyxl
    import openpyxl
    wb = openpyxl.load_workbook(out_xlsx)
    sheet_names = wb.sheetnames

    assert "Piping Engineering" in sheet_names
    assert "Electrical Engineering" in sheet_names
    assert "GPD" in sheet_names
    assert "Plakon" in sheet_names

    piping_ws = wb["Piping Engineering"]
    # Row 5 should have Piping comment
    assert piping_ws.cell(row=5, column=8).value == "Piping clash with tray"

    elec_ws = wb["Electrical Engineering"]
    # Row 5 should have Electrical comment
    assert elec_ws.cell(row=5, column=8).value == "Electrical conduit ungrounded"

    gpd_ws = wb["GPD"]
    # GPD has zero comments — should have standard empty header rows
    assert gpd_ws.cell(row=5, column=8).value in ("", None)


def test_analytics_department_category_counts(test_db_engine: DatabaseEngine, sample_pdf_drawing: Path):
    pdf_service = PDFService(pdf_loader=PyMuPDFAdapter())
    doc_dto = pdf_service.process_pdf_document(sample_pdf_drawing)
    comment_repo = CommentRepository(test_db_engine)
    dept_repo = EngineeringDepartmentRepository(test_db_engine)
    dwg_repo = DrawingRepository(test_db_engine)
    analytics_service = AnalyticsService(test_db_engine)

    piping_dept = dept_repo.get_or_create_department("Piping Engineering")
    dwg_rec = dwg_repo.save_drawing_from_dto(doc_dto, department_id=piping_dept["id"])

    comment_repo.save_comment(
        drawing_id=dwg_rec["id"],
        page_number=1,
        raw_text="Dimension missing",
        bbox=(0, 0, 10, 10),
        category_name="Dimensional",
        department_id=piping_dept["id"],
    )

    counts = comment_repo.get_department_category_counts(dwg_rec["id"])
    assert ("Piping Engineering", "Dimensional") in counts
    assert counts[("Piping Engineering", "Dimensional")] == 1

    pareto = analytics_service.get_pareto_analysis(department_name="Piping Engineering")
    assert len(pareto) >= 1
    assert pareto[0].category_name == "Dimensional"
    assert pareto[0].count == 1

